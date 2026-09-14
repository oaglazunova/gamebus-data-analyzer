from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _normalize_email(
    value: Any,
) -> str:
    if value is None:
        return ""

    return str(value).strip().casefold()


def _normalize_header(
    value: Any,
) -> str:
    if value is None:
        return ""

    return (
        str(value)
        .strip()
        .casefold()
        .replace(" ", "_")
        .replace("-", "_")
    )


class CredentialStore:
    """
    In-memory participant credentials.

    Passwords are deliberately private and are not included
    in repr(), manifests, UI tables, or exported structures.
    """

    __slots__ = (
        "_passwords_by_email",
        "source_filename",
    )

    def __init__(
        self,
        *,
        passwords_by_email: dict[str, str],
        source_filename: str,
    ) -> None:
        self._passwords_by_email = dict(
            passwords_by_email
        )

        self.source_filename = (
            source_filename
        )

    @property
    def count(self) -> int:
        return len(
            self._passwords_by_email
        )

    @property
    def emails(self) -> frozenset[str]:
        return frozenset(
            self._passwords_by_email
        )

    def has_email(
        self,
        email: str | None,
    ) -> bool:
        normalized = _normalize_email(
            email
        )

        if not normalized:
            return False

        return (
            normalized
            in self._passwords_by_email
        )

    def get_password(
        self,
        email: str,
    ) -> str:
        normalized = _normalize_email(
            email
        )

        if not normalized:
            raise KeyError(
                "Participant email is empty."
            )

        return self._passwords_by_email[
            normalized
        ]

    def __repr__(self) -> str:
        return (
            "CredentialStore("
            f"source_filename="
            f"{self.source_filename!r}, "
            f"count={self.count}"
            ")"
        )


@dataclass(
    frozen=True
)
class CohortMatchSummary:
    campaign_accounts: int
    credentials_accounts: int
    matched_accounts: int
    credentials_not_in_campaign: int


def load_participant_credentials(
    path: Path,
) -> CredentialStore:
    """
    Load participant email/password credentials from XLSX.

    The workbook may contain additional columns; they are
    ignored.

    A worksheet is usable when its first row contains both:

        email
        password

    Header matching is case-insensitive.
    """
    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(
            path
        )

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        selected_sheet = None
        email_index = None
        password_index = None

        for sheet in workbook.worksheets:
            first_row = next(
                sheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                None,
            )

            if first_row is None:
                continue

            normalized_headers = [
                _normalize_header(
                    value
                )
                for value in first_row
            ]

            if (
                "email"
                not in normalized_headers
                or "password"
                not in normalized_headers
            ):
                continue

            selected_sheet = sheet

            email_index = (
                normalized_headers.index(
                    "email"
                )
            )

            password_index = (
                normalized_headers.index(
                    "password"
                )
            )

            break

        if selected_sheet is None:
            raise ValueError(
                "Credentials workbook does not "
                "contain a worksheet with both "
                "'email' and 'password' columns."
            )

        passwords_by_email: dict[
            str,
            str,
        ] = {}

        for row_number, row in enumerate(
            selected_sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            email_value = (
                row[email_index]
                if email_index < len(row)
                else None
            )

            password_value = (
                row[password_index]
                if password_index < len(row)
                else None
            )

            email = _normalize_email(
                email_value
            )

            password = (
                ""
                if password_value is None
                else str(password_value)
            )

            # Ignore fully empty rows.
            if (
                not email
                and not password
            ):
                continue

            if not email:
                raise ValueError(
                    "Credentials workbook contains "
                    f"a row without email: "
                    f"row {row_number}."
                )

            if not password:
                raise ValueError(
                    "Credentials workbook contains "
                    "an account without password: "
                    f"row {row_number}."
                )

            if (
                email
                in passwords_by_email
            ):
                raise ValueError(
                    "Credentials workbook contains "
                    "duplicate participant email: "
                    f"{email}"
                )

            passwords_by_email[
                email
            ] = password

        return CredentialStore(
            passwords_by_email=(
                passwords_by_email
            ),
            source_filename=path.name,
        )

    finally:
        workbook.close()


def build_initial_cohort_candidates(
    *,
    campaign_users: list[
        dict[str, str | None]
    ],
    credentials: CredentialStore | None,
) -> tuple[
    list[dict[str, Any]],
    CohortMatchSummary,
]:
    """
    Build safe UI-ready cohort candidates.

    Without credentials:
        all campaign accounts are selected for analysis.

    With credentials:
        only matching accounts are initially selected.

    Participant-level extraction is initially selected only
    when the account is selected for analysis AND credentials
    are available.
    """
    candidates: list[
        dict[str, Any]
    ] = []

    campaign_emails: set[str] = set()

    matched_accounts = 0

    for user in campaign_users:
        email = user.get(
            "email"
        )

        normalized_email = (
            _normalize_email(
                email
            )
        )

        if normalized_email:
            campaign_emails.add(
                normalized_email
            )

        credentials_available = (
            credentials is not None
            and credentials.has_email(
                email
            )
        )

        if (
            credentials is None
        ):
            selected_for_analysis = True
        else:
            selected_for_analysis = (
                credentials_available
            )

        selected_for_extraction = (
            selected_for_analysis
            and credentials_available
        )

        if credentials_available:
            matched_accounts += 1

        candidates.append(
            {
                "account_id": user.get(
                    "account_id"
                ),
                "pid": user.get(
                    "pid"
                ),
                "email": email,
                "credentials_available": (
                    credentials_available
                ),
                "selected_for_analysis": (
                    selected_for_analysis
                ),
                "selected_for_participant_extraction": (
                    selected_for_extraction
                ),
            }
        )

    credential_emails = (
        set()
        if credentials is None
        else set(
            credentials.emails
        )
    )

    summary = CohortMatchSummary(
        campaign_accounts=len(
            campaign_users
        ),
        credentials_accounts=len(
            credential_emails
        ),
        matched_accounts=(
            matched_accounts
        ),
        credentials_not_in_campaign=len(
            credential_emails
            - campaign_emails
        ),
    )

    return (
        candidates,
        summary,
    )